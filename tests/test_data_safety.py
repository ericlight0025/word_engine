from __future__ import annotations

import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core.csv_reader import read_csv
from core.data_writer import write_dataset
from core.doc_converter import prepare_template
from core.excel_reader import read_excel
from core.settings_store import AppSettings, SettingsStore


class DataSafetyTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def make_workbook(self):
        path = self.root / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "名單"
        sheet.append(["姓名", "金額", "計算"])
        sheet.append(["甲", 100, "=B2*2"])
        sheet.append([None, None, None])
        sheet.append(["乙", 200, "=B4*2"])
        sheet["A4"].font = Font(bold=True, color="FF0000")
        sheet["B4"].number_format = "#,##0.00"
        sheet.column_dimensions["A"].width = 24
        sheet.freeze_panes = "A2"
        other = workbook.create_sheet("其他資料")
        other["A1"] = "重要資料"
        other["B1"] = "=1+2"
        workbook.save(path)
        workbook.close()
        return path

    def test_edit_preserves_sheets_formulas_styles_and_original_row_positions(self):
        path = self.make_workbook()
        dataset = read_excel(path)
        self.assertEqual(dataset.row_numbers, [2, 4])
        dataset.rows[1]["姓名"] = "乙修改"
        write_dataset(path, dataset.headers, dataset.rows, source_dataset=dataset)
        workbook = load_workbook(path, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["名單"]
        self.assertEqual(workbook.sheetnames, ["名單", "其他資料"])
        self.assertEqual(sheet["A4"].value, "乙修改")
        self.assertIsNone(sheet["A3"].value)
        self.assertEqual(sheet["C2"].value, "=B2*2")
        self.assertEqual(sheet["C4"].value, "=B4*2")
        self.assertEqual(sheet["B4"].data_type, "n")
        self.assertTrue(sheet["A4"].font.bold)
        self.assertEqual(sheet["B4"].number_format, "#,##0.00")
        self.assertEqual(sheet.column_dimensions["A"].width, 24)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(workbook["其他資料"]["B1"].value, "=1+2")

    def test_edited_excel_values_starting_with_equals_remain_text(self):
        path = self.make_workbook()
        dataset = read_excel(path)
        dataset.rows[0]["姓名"] = "=1+1"
        write_dataset(path, dataset.headers, dataset.rows, source_dataset=dataset)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["A2"].value, "=1+1")
        self.assertEqual(workbook.active["A2"].data_type, "s")

    def test_failed_excel_publish_keeps_original_bytes(self):
        path = self.make_workbook()
        original = path.read_bytes()
        dataset = read_excel(path)
        dataset.rows[0]["姓名"] = "修改"
        with patch("core.file_store.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                write_dataset(path, dataset.headers, dataset.rows, source_dataset=dataset)
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(list(self.root.glob(".word-engine-*")), [])

    def test_external_excel_modification_cannot_be_overwritten(self):
        path = self.make_workbook()
        dataset = read_excel(path)
        workbook = load_workbook(path)
        workbook.active["A2"] = "外部修改"
        workbook.save(path)
        workbook.close()
        original = path.read_bytes()
        dataset.rows[0]["姓名"] = "畫面修改"
        with self.assertRaisesRegex(ValueError, "其他程式修改"):
            write_dataset(path, dataset.headers, dataset.rows, source_dataset=dataset)
        self.assertEqual(path.read_bytes(), original)

    def test_csv_duplicate_blank_headers_and_extra_values_are_rejected(self):
        path = self.root / "source.csv"
        for content in ("姓名,姓名\n甲,乙\n", "姓名,\n甲,乙\n", "姓名\n甲,乙\n", "姓名, 姓名 \n甲,乙\n"):
            with self.subTest(content=content):
                path.write_text(content, encoding="utf-8-sig")
                with self.assertRaises(ValueError):
                    read_csv(path)

    def test_csv_normalizes_headers_and_pads_missing_fields(self):
        path = self.root / "source.csv"
        path.write_text(' 姓名 ,日期\n"甲,乙"\n\n', encoding="utf-8-sig")
        dataset = read_csv(path)
        self.assertEqual(dataset.headers, ["姓名", "日期"])
        self.assertEqual(dataset.rows, [{"姓名": "甲,乙", "日期": ""}])

    def test_failed_csv_write_keeps_source(self):
        path = self.root / "source.csv"
        path.write_text("姓名\n甲\n", encoding="utf-8-sig")
        original = path.read_bytes()
        with patch("core.file_store.os.fsync", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                write_dataset(path, ["姓名"], [{"姓名": "修改"}])
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(list(self.root.glob(".word-engine-*")), [])

    def test_external_csv_modification_cannot_be_overwritten(self):
        path = self.root / "source.csv"
        path.write_text("姓名\n甲\n", encoding="utf-8-sig")
        dataset = read_csv(path)
        path.write_text("姓名\n外部修改\n", encoding="utf-8-sig")
        with self.assertRaisesRegex(ValueError, "其他程式修改"):
            write_dataset(path, dataset.headers, dataset.rows, source_dataset=dataset)
        self.assertIn("外部修改", path.read_text(encoding="utf-8-sig"))

    def test_extra_dataset_keys_are_not_silently_discarded(self):
        with self.assertRaisesRegex(ValueError, "未列入表頭"):
            write_dataset(self.root / "out.xlsx", ["姓名"], [{"姓名": "甲", "秘密": "乙"}])
        self.assertFalse((self.root / "out.xlsx").exists())

    def test_excel_handles_close_on_header_validation_failure(self):
        workbook = Mock()
        workbook.active.iter_rows.return_value = iter([("姓名", "姓名"), ("甲", "乙")])
        with patch("core.excel_reader.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "重複"):
                read_excel(self.root / "invalid.xlsx")
        workbook.close.assert_called_once()

    def test_xls_is_rejected_with_conversion_guidance(self):
        with self.assertRaisesRegex(ValueError, "另存為 .xlsx"):
            read_excel(self.root / "legacy.xls")

    def test_invalid_settings_types_do_not_crash_startup(self):
        store = SettingsStore(self.root / "settings.json")
        for payload in ([], None, 123, {"font_scale": "broken"}, {"font_scale": None}, {"font_scale": -50}):
            with self.subTest(payload=payload):
                store.path.write_text(json.dumps(payload), encoding="utf-8")
                settings = store.load()
                self.assertGreaterEqual(settings.font_scale, 90)
                self.assertLessEqual(settings.font_scale, 130)

    def test_settings_save_failure_keeps_valid_json(self):
        store = SettingsStore(self.root / "settings.json")
        store.save(AppSettings(data_dir="old"))
        original = store.path.read_bytes()
        with patch("core.file_store.os.replace", side_effect=OSError("locked")):
            with self.assertRaises(OSError):
                store.save(AppSettings(data_dir="new"))
        self.assertEqual(store.path.read_bytes(), original)
        self.assertEqual(store.load().data_dir, "old")

    def test_converter_cleans_up_after_timeout_and_uses_isolated_profile(self):
        source = self.root / "legacy.doc"
        source.write_bytes(b"original")
        directories = []

        def make_temporary(**kwargs):
            result = tempfile.TemporaryDirectory(dir=self.root, **kwargs)
            directories.append(Path(result.name))
            return result

        with (
            patch("core.doc_converter.libreoffice_exists", return_value=True),
            patch("core.doc_converter.TemporaryDirectory", side_effect=make_temporary),
            patch("core.doc_converter.subprocess.run", side_effect=subprocess.TimeoutExpired("soffice", 60)) as run,
        ):
            with self.assertRaisesRegex(RuntimeError, "超過 60 秒"):
                prepare_template(source)
        self.assertEqual(run.call_args.kwargs["timeout"], 60)
        self.assertTrue(any(arg.startswith("-env:UserInstallation=file:") for arg in run.call_args.args[0]))
        self.assertFalse(directories[0].exists())
        self.assertEqual(source.read_bytes(), b"original")


if __name__ == "__main__":
    unittest.main()
