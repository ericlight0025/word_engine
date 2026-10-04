from __future__ import annotations

import tempfile
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from docx import Document

from core.template_engine import extract_tags, extract_template_preview, merge_documents, sanitize_filename


class MergeSafetyTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.template = self.root / "template.docx"
        document = Document()
        document.add_paragraph("姓名：{{ 姓名 }}")
        document.save(self.template)

    def test_xml_special_characters_remain_exact_text(self):
        value = '甲 & 乙 <丙> "丁"'
        summary = merge_documents(self.template, [{"姓名": value}], self.root / "out", "")
        self.assertEqual(summary.failure_count, 0)
        self.assertEqual(Document(summary.output_files[0]).paragraphs[0].text, "姓名：" + value)
        with zipfile.ZipFile(summary.output_files[0]) as archive:
            for name in archive.namelist():
                if name.endswith(".xml"):
                    ElementTree.fromstring(archive.read(name))

    def test_duplicate_names_preserve_each_document_and_existing_file(self):
        output = self.root / "out"
        output.mkdir()
        (output / "same.docx").write_bytes(b"existing document")
        summary = merge_documents(self.template, [{"姓名": "甲", "編號": "same"}, {"姓名": "乙", "編號": "same"}], output, "編號")
        self.assertEqual(summary.success_count, 2)
        self.assertEqual([p.name for p in summary.output_files], ["same_1.docx", "same_2.docx"])
        self.assertEqual((output / "same.docx").read_bytes(), b"existing document")
        self.assertEqual([Document(p).paragraphs[0].text for p in summary.output_files], ["姓名：甲", "姓名：乙"])

    def test_output_matching_template_name_keeps_template_unchanged(self):
        original = self.template.read_bytes()
        summary = merge_documents(self.template, [{"姓名": "甲", "編號": "template"}], self.root, "編號")
        self.assertEqual(summary.output_files[0].name, "template_1.docx")
        self.assertEqual(self.template.read_bytes(), original)

    def test_concurrent_batches_keep_all_complete_outputs(self):
        def merge(index):
            return merge_documents(self.template, [{"姓名": str(index), "編號": "same"}], self.root / "out", "編號")
        with ThreadPoolExecutor(max_workers=4) as pool:
            summaries = list(pool.map(merge, range(4)))
        paths = [result.output_files[0] for result in summaries]
        self.assertEqual(len(set(paths)), 4)
        self.assertEqual({Document(p).paragraphs[0].text for p in paths}, {f"姓名：{index}" for index in range(4)})

    def test_failed_publish_keeps_existing_output_and_cleans_temp(self):
        output = self.root / "out"
        output.mkdir()
        original = output / "same.docx"
        original.write_bytes(b"old")
        with patch("core.file_store.os.link", side_effect=OSError("disk error")):
            summary = merge_documents(self.template, [{"姓名": "甲", "編號": "same"}], output, "編號")
        self.assertEqual(summary.failure_count, 1)
        self.assertEqual(summary.success_count, 0)
        self.assertEqual(original.read_bytes(), b"old")
        self.assertEqual(list(output.glob(".word-engine-*")), [])

    def test_split_runs_filters_conditions_and_footer_use_real_variables(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("{{ 姓")
        paragraph.add_run("名 | upper }}")
        document.add_paragraph("{% if 公司 %}{{ 公司 }}{% endif %}")
        document.sections[0].footer.paragraphs[0].text = "{{ 日期 }}"
        document.save(self.template)
        self.assertEqual(extract_tags(self.template), ["公司", "姓名", "日期"])
        self.assertIn("{{ 姓名 | upper }}", extract_template_preview(self.template))

    def test_template_preview_includes_body_table_header_footer_and_entities(self):
        document = Document()
        document.add_paragraph("內容 A & B")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "表格內容"
        document.sections[0].header.paragraphs[0].text = "頁首"
        document.sections[0].footer.paragraphs[0].text = "頁尾"
        document.save(self.template)
        text = extract_template_preview(self.template)
        for value in ("內容 A & B", "表格內容", "頁首", "頁尾"):
            self.assertIn(value, text)
        self.assertNotIn("<w:", text)

    def test_missing_fields_and_name_fallback_count_unique_warned_rows(self):
        summary = merge_documents(self.template, [{"編號": ""}], self.root / "out", "編號")
        self.assertEqual(summary.success_count, 1)
        self.assertEqual(summary.warning_count, 1)
        self.assertEqual(len(summary.warnings), 2)
        self.assertEqual(Document(summary.output_files[0]).paragraphs[0].text, "姓名：")

    def test_invalid_row_does_not_abort_later_valid_rows(self):
        summary = merge_documents(self.template, [{"姓名": "bad\x01"}, {"姓名": "正確"}], self.root / "out", "")
        self.assertEqual((summary.failure_count, summary.success_count), (1, 1))
        self.assertEqual(summary.failures[0].row_index, 1)
        self.assertEqual(Document(summary.output_files[0]).paragraphs[0].text, "姓名：正確")

    def test_windows_reserved_names_control_characters_and_long_unicode_names(self):
        for name in ("CON", "con.txt", "LPT1", "COM¹", "NUL"):
            with self.subTest(name=name):
                self.assertTrue(sanitize_filename(name).startswith("_"))
        self.assertNotIn("\x00", sanitize_filename("bad\x00name"))
        self.assertLessEqual(len(sanitize_filename("漢" * 200).encode("utf-8")), 180)


if __name__ == "__main__":
    unittest.main()
